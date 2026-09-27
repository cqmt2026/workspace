# -*- coding: utf-8 -*-
"""
生成机械行业完整拨打清单（Excel格式）
"""
import sys, sqlite3, json
from datetime import datetime
sys.stdout.reconfigure(encoding='utf-8')

conn = sqlite3.connect('data/baidu_crm.db')
conn.row_factory = sqlite3.Row
c = conn.cursor()

print('='*70)
print('[Mechanical Industry Full Call List Generation]')
print('='*70)
print(f'Generation time: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}')
print('='*70)

# Query P0-P2 customers with phone, ordered by priority
print('\n[Step 1] Fetching P0-P2 customers with phone...')
c.execute('''
    SELECT id, name, industry, district, district_priority, phone, address, lat, lng
    FROM customers
    WHERE (industry LIKE '%机械%' OR industry = '液压设备' 
           OR industry = '工程机械' OR industry = '数控机床'
           OR industry = '精密机械' OR industry = '自动化设备')
      AND district_priority IN ('P0', 'P1', 'P2')
      AND phone IS NOT NULL
      AND phone != ''
    ORDER BY 
        CASE district_priority 
            WHEN 'P0' THEN 1 
            WHEN 'P1' THEN 2 
            WHEN 'P2' THEN 3 
        END,
        district,
        industry
''')
rows = c.fetchall()
print(f'  Total: {len(rows)}')

# Group by priority
p0_list = [r for r in rows if r['district_priority'] == 'P0']
p1_list = [r for r in rows if r['district_priority'] == 'P1']
p2_list = [r for r in rows if r['district_priority'] == 'P2']

print(f'  P0: {len(p0_list)}')
print(f'  P1: {len(p1_list)}')
print(f'  P2: {len(p2_list)}')

# Generate Excel file
print('\n[Step 2] Generating Excel file...')
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

wb = openpyxl.Workbook()

# Sheet 1: P0 Main Urban
ws_p0 = wb.active
ws_p0.title = 'P0_主城区'

# Header
headers = ['序号', '企业ID', '企业名称', '行业', '区域', '电话', '地址', '纬度', '经度', '优先级', '话术建议']
ws_p0.append(headers)
header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
header_font = Font(bold=True, color='FFFFFF')
for cell in ws_p0[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center', vertical='center')

# Data
for i, r in enumerate(p0_list, 1):
    ws_p0.append([
        i,
        r['id'],
        r['name'],
        r['industry'],
        r['district'] or 'N/A',
        r['phone'],
        r['address'] or '',
        r['lat'] or '',
        r['lng'] or '',
        'P0',
        f'重庆{r["industry"]}百度几乎无投放，{r["district"]}区域蓝海'
    ])

# Set column widths
ws_p0.column_dimensions['A'].width = 8
ws_p0.column_dimensions['B'].width = 10
ws_p0.column_dimensions['C'].width = 40
ws_p0.column_dimensions['D'].width = 15
ws_p0.column_dimensions['E'].width = 12
ws_p0.column_dimensions['F'].width = 20
ws_p0.column_dimensions['G'].width = 30
ws_p0.column_dimensions['H'].width = 12
ws_p0.column_dimensions['I'].width = 12
ws_p0.column_dimensions['J'].width = 10
ws_p0.column_dimensions['K'].width = 35

# Sheet 2: P1 Suburban
ws_p1 = wb.create_sheet('P1_近郊')
ws_p1.append(headers)
for cell in ws_p1[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center', vertical='center')

for i, r in enumerate(p1_list, 1):
    ws_p1.append([
        i,
        r['id'],
        r['name'],
        r['industry'],
        r['district'] or 'N/A',
        r['phone'],
        r['address'] or '',
        r['lat'] or '',
        r['lng'] or '',
        'P1',
        f'重庆{r["industry"]}百度几乎无投放，{r["district"]}区域蓝海'
    ])

for col in 'ABCDEFGHIJK':
    ws_p1.column_dimensions[col].width = ws_p0.column_dimensions[col].width

# Sheet 3: P2 Other Districts
ws_p2 = wb.create_sheet('P2_其他区县')
ws_p2.append(headers)
for cell in ws_p2[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center', vertical='center')

for i, r in enumerate(p2_list, 1):
    ws_p2.append([
        i,
        r['id'],
        r['name'],
        r['industry'],
        r['district'] or 'N/A',
        r['phone'],
        r['address'] or '',
        r['lat'] or '',
        r['lng'] or '',
        'P2',
        f'重庆{r["industry"]}百度几乎无投放，{r["district"]}区域蓝海'
    ])

for col in 'ABCDEFGHIJK':
    ws_p2.column_dimensions[col].width = ws_p0.column_dimensions[col].width

# Sheet 4: Summary
ws_summary = wb.create_sheet('汇总统计', 0)
ws_summary.append(['机械行业拨打清单汇总'])
ws_summary.append([])
ws_summary.append(['生成时间', datetime.now().strftime('%Y-%m-%d %H:%M:%S')])
ws_summary.append(['数据来源', '百度POI补抓 + 高德POI'])
ws_summary.append(['优先级规则', 'P0(主城区) > P1(近郊) > P2(其他区县)'])
ws_summary.append([])
ws_summary.append(['优先级', '区域范围', '数量', '占比'])
ws_summary.append(['P0', '主城区9区+高新区+两江新区', len(p0_list), f'{len(p0_list)/len(rows)*100:.1f}%'])
ws_summary.append(['P1', '璧山/江津/合川/永川/大足/铜梁', len(p1_list), f'{len(p1_list)/len(rows)*100:.1f}%'])
ws_summary.append(['P2', '綦江/南川/长寿/荣昌等', len(p2_list), f'{len(p2_list)/len(rows)*100:.1f}%'])
ws_summary.append(['总计', '-', len(rows), '100%'])
ws_summary.append([])
ws_summary.append(['拨打建议'])
ws_summary.append(['1. 优先拨打P0：主城区企业决策快，百度投放意愿高'])
ws_summary.append(['2. 话术核心：先问"你们在百度能搜到吗"，制造焦虑'])
ws_summary.append(['3. 算账策略：B2B制造业，投1万换10万订单'])
ws_summary.append(['4. 记录要点：痛点+当前获客方式+预算量级+决策人'])

for col in ws_summary.columns:
    ws_summary.column_dimensions[col[0].column_letter].width = 25

# Save
output_file = f'assets/机械行业完整拨打清单_P0P1P2_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
wb.save(output_file)
print(f'\n💾 Excel saved: {output_file}')

# Also save JSON
json_file = f'assets/机械行业拨打清单_P0P1P2_{datetime.now().strftime("%Y%m%d")}.json'
with open(json_file, 'w', encoding='utf-8') as f:
    json.dump({
        'generation_time': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
        'total': len(rows),
        'p0_count': len(p0_list),
        'p1_count': len(p1_list),
        'p2_count': len(p2_list),
        'customers': [
            {
                'seq': i,
                'priority': 'P0',
                'id': r['id'],
                'name': r['name'],
                'industry': r['industry'],
                'district': r['district'],
                'phone': r['phone'],
                'address': r['address'],
                'lat': r['lat'],
                'lng': r['lng'],
                'script': f'重庆{r["industry"]}百度几乎无投放，{r["district"]}区域蓝海'
            }
            for i, r in enumerate(p0_list, 1)
        ] + [
            {
                'seq': i,
                'priority': 'P1',
                'id': r['id'],
                'name': r['name'],
                'industry': r['industry'],
                'district': r['district'],
                'phone': r['phone'],
                'address': r['address'],
                'lat': r['lat'],
                'lng': r['lng'],
                'script': f'重庆{r["industry"]}百度几乎无投放，{r["district"]}区域蓝海'
            }
            for i, r in enumerate(p1_list, 1)
        ] + [
            {
                'seq': i,
                'priority': 'P2',
                'id': r['id'],
                'name': r['name'],
                'industry': r['industry'],
                'district': r['district'],
                'phone': r['phone'],
                'address': r['address'],
                'lat': r['lat'],
                'lng': r['lng'],
                'script': f'重庆{r["industry"]}百度几乎无投放，{r["district"]}区域蓝海'
            }
            for i, r in enumerate(p2_list, 1)
        ]
    }, f, ensure_ascii=False, indent=2)
print(f'💾 JSON saved: {json_file}')

print('\n' + '='*70)
print('[Generation Complete]')
print('='*70)
print(f'Total: {len(rows)} customers')
print(f'P0: {len(p0_list)}')
print(f'P1: {len(p1_list)}')
print(f'P2: {len(p2_list)}')
print(f'\nFiles:')
print(f'  - Excel: {output_file}')
print(f'  - JSON: {json_file}')

conn.close()
