# -*- coding: utf-8 -*-
"""解析经信委创新型中小企业名单"""
import openpyxl, sys, csv, sqlite3, os
sys.stdout.reconfigure(encoding='utf-8')

BASE = r'C:\Users\Administrator\.openclaw\workspace\workspace-baiduagent\data\source\jingxinwei'
MASTER_DIR = r'C:\Users\Administrator\.openclaw\workspace\workspace-baiduagent\data\master'

print("=" * 70)
print("【经信委创新型中小企业名单 - 解析入库】")
print("=" * 70)

all_companies = []
for fname in ['2025_创新型_申报通过.xlsx', '2025_创新型_复核通过.xlsx']:
    fpath = os.path.join(BASE, fname)
    if not os.path.exists(fpath):
        continue
    wb = openpyxl.load_workbook(fpath, data_only=True)
    ws = wb.active
    print(f"\n📄 {fname}: {ws.max_row}行, {ws.max_column}列")
    print(f"   表头: {[ws.cell(1,c).value for c in range(1, min(ws.max_column+1,5))]}")
    
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        name = str(row[0]).strip() if row[0] else ''
        if not name or name == 'None':
            continue
        # 其他列
        code = str(row[1]).strip() if len(row) > 1 and row[1] else ''
        dist = str(row[2]).strip() if len(row) > 2 and row[2] else ''
        
        all_companies.append({
            'name': name,
            'code': code,
            'district': dist,
            'source': '经信委创新型中小企业2025'
        })
        if len(all_companies) <= 3:
            print(f"   样例: {name}, 区={dist}")

print(f"\n✅ 共解析 {len(all_companies)} 家企业")

# 保存到master
out_path = os.path.join(MASTER_DIR, '创新型中小企业_jingxinwei_2025.csv')
with open(out_path, 'w', encoding='utf-8-sig', newline='') as f:
    writer = csv.DictWriter(f, fieldnames=['name','code','district','source'])
    writer.writeheader()
    for c in all_companies:
        writer.writerow(c)

print(f"✅ 已保存到: {out_path}")

# 统计各区数量
districts = {}
for c in all_companies:
    d = c['district'] or '未知'
    districts[d] = districts.get(d, 0) + 1
print(f"\n📊 各区分布:")
for d, cnt in sorted(districts.items(), key=lambda x: -x[1]):
    print(f"   {d}: {cnt}家")
