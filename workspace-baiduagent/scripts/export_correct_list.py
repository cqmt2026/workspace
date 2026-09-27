#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""导出下周三天拨打清单（正确日期）"""
import sqlite3
import random
import os

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
except ImportError:
    print('请安装openpyxl: pip install openpyxl')
    exit(1)

DB = 'data/baidu_crm.db'
conn = sqlite3.connect(DB)
c = conn.cursor()

# 获取所有有电话的企业，排除宠物医院
c.execute('''SELECT id, name, industry, phone, district 
             FROM customers 
             WHERE phone IS NOT NULL 
             AND phone != ""
             AND industry != "宠物医院"
             ORDER BY id''')
all_customers = c.fetchall()

print(f'总客户数（排除宠物医院）：{len(all_customers)}家')

# 打乱
random.seed(20260927)
random.shuffle(all_customers)

# 分组 - 下周一到周三：09-28, 09-29, 09-30
day1 = all_customers[:200]   # 周一 09-28
day2 = all_customers[200:400] # 周二 09-29
day3 = all_customers[400:600] # 周三 09-30

print(f'周一 09-28：{len(day1)}家')
print(f'周二 09-29：{len(day2)}家')
print(f'周三 09-30：{len(day3)}家')

# 按行业统计
industry_count = {}
for day_name, day_data in [('周一', day1), ('周二', day2), ('周三', day3)]:
    for cust in day_data:
        industry = cust[2]
        industry_count[industry] = industry_count.get(industry, 0) + 1

print('\n行业分布TOP10:')
for ind, cnt in sorted(industry_count.items(), key=lambda x: -x[1])[:10]:
    print(f'  {ind}: {cnt}家')

# 生成Markdown表格
lines = []
lines.append('# 下周三天拨打清单')
lines.append('')
lines.append('> 生成时间：2026-09-27')
lines.append('> 执行日期：周一09-28 / 周二09-29 / 周三09-30')
lines.append('> 每天200家，共600家，不重复')
lines.append('> 已过滤：宠物医院（上周验证零需求）')
lines.append('')

for day_name, day_data in [('周一 09-28', day1), ('周二 09-29', day2), ('周三 09-30', day3)]:
    lines.append(f'## {day_name}')
    lines.append('')
    lines.append('| 序号 | 企业名称 | 行业 | 电话 | 区域 |')
    lines.append('|------|---------|------|------|------|')
    
    for i, cust in enumerate(day_data, 1):
        idx, name, industry, phone, district = cust
        lines.append(f'| {i} | {name} | {industry} | {phone} | {district} |')
    
    lines.append('')
    lines.append('---')
    lines.append('')

# 写入Markdown
md_path = 'assets/下周三天拨打清单_0928-0930.md'
with open(md_path, 'w', encoding='utf-8') as f:
    f.write('\n'.join(lines))

# 创建Excel
wb = openpyxl.Workbook()

# 样式
header_font = Font(bold=True, size=11, color="FFFFFF")
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
thin_border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

def create_sheet(ws, title, data):
    ws.title = title
    ws.append([title])
    ws.append([])
    ws.append(['序号', '企业名称', '行业', '电话', '区域', '拨打状态', '意向等级', '备注'])
    
    for i, cust in enumerate(data, 1):
        idx, name, industry, phone, district = cust
        ws.append([i, name, industry, phone, district, '', '', ''])
    
    # 设置表头样式
    for cell in ws[3]:
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal='center', vertical='center')
    
    # 调整列宽
    ws.column_dimensions['A'].width = 6
    ws.column_dimensions['B'].width = 35
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 18
    ws.column_dimensions['E'].width = 12
    ws.column_dimensions['F'].width = 12
    ws.column_dimensions['G'].width = 12
    ws.column_dimensions['H'].width = 20

ws1 = wb.create_sheet()
ws2 = wb.create_sheet()
ws3 = wb.create_sheet()

create_sheet(ws1, '周一 09-28', day1)
create_sheet(ws2, '周二 09-29', day2)
create_sheet(ws3, '周三 09-30', day3)

# 添加汇总页
ws_summary = wb.create_sheet('汇总', 0)
ws_summary.append(['下周三天拨打清单汇总'])
ws_summary.append([])
ws_summary.append(['日期', '行业分布TOP5', '目标数量', '实际拨打', '意向数', '转化率'])
ws_summary.append(['周一 09-28', '', '200', '', '', ''])
ws_summary.append(['周二 09-29', '', '200', '', '', ''])
ws_summary.append(['周三 09-30', '', '200', '', '', ''])
ws_summary.append(['总计', '', '600', '', '', ''])

output_path = 'assets/下周三天拨打清单_0928-0930.xlsx'
wb.save(output_path)
conn.close()

print(f'\n已导出Markdown：{md_path}')
print(f'已导出Excel：{output_path}')
