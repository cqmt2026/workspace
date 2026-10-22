#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""生成《路上有惊慌》好书推荐PPT - 小学四年级角度"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# 输出路径
OUTPUT_DIR = r"C:\Users\Administrator\.openclaw\workspace\output"
os.makedirs(OUTPUT_DIR, exist_ok=True)
OUTPUT_PATH = os.path.join(OUTPUT_DIR, "路上有惊慌_好书推荐.pptx")

# 创建演示文稿 (16:9)
prs = Presentation()
prs.slide_width = Inches(13.33)
prs.slide_height = Inches(7.5)

# 颜色方案 - 温暖活泼
COLORS = {
    'primary': RGBColor(255, 140, 0),    # 橙色
    'secondary': RGBColor(255, 193, 7),  # 黄色
    'accent': RGBColor(76, 175, 80),     # 绿色
    'text': RGBColor(51, 51, 51),        # 深灰
    'light_bg': RGBColor(255, 248, 220), # 浅黄背景
}

def add_title_slide(prs, title, subtitle):
    """标题页"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # 空白布局
    
    # 背景
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLORS['light_bg']
    bg.line.fill.background()
    
    # 标题
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(12.33), Inches(1.5))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(54)
    p.font.bold = True
    p.font.color.rgb = COLORS['primary']
    p.alignment = PP_ALIGN.CENTER
    
    # 副标题
    sub_box = slide.shapes.add_textbox(Inches(0.5), Inches(4.2), Inches(12.33), Inches(1))
    tf = sub_box.text_frame
    p = tf.paragraphs[0]
    p.text = subtitle
    p.font.size = Pt(28)
    p.font.color.rgb = COLORS['text']
    p.alignment = PP_ALIGN.CENTER
    
    return slide

def add_content_slide(prs, title, content_list, highlight=None):
    """内容页"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    # 标题栏
    title_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.2))
    title_bar.fill.solid()
    title_bar.fill.fore_color.rgb = COLORS['primary']
    title_bar.line.fill.background()
    
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.33), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    
    # 内容
    content_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.73), Inches(5.5))
    tf = content_box.text_frame
    tf.word_wrap = True
    
    for i, item in enumerate(content_list):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        
        p.text = "• " + item
        p.font.size = Pt(24)
        p.font.color.rgb = COLORS['text']
        p.space_before = Pt(16)
        p.line_spacing = 1.5
        
        if highlight and i == highlight:
            p.font.bold = True
            p.font.color.rgb = COLORS['accent']
    
    return slide

def add_quote_slide(prs, quote, source):
    """引用页"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    # 背景
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLORS['secondary']
    bg.line.fill.background()
    
    # 引用文字
    quote_box = slide.shapes.add_textbox(Inches(1.5), Inches(2), Inches(10.33), Inches(3))
    tf = quote_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = f'"{quote}"'
    p.font.size = Pt(36)
    p.font.italic = True
    p.font.color.rgb = COLORS['text']
    p.alignment = PP_ALIGN.CENTER
    p.space_before = Pt(20)
    
    # 来源
    source_box = slide.shapes.add_textbox(Inches(1.5), Inches(5.2), Inches(10.33), Inches(0.8))
    tf = source_box.text_frame
    p = tf.paragraphs[0]
    p.text = f"—— {source}"
    p.font.size = Pt(24)
    p.font.color.rgb = COLORS['text']
    p.alignment = PP_ALIGN.RIGHT
    
    return slide

# ===== 制作5页PPT =====

# 第1页：封面
add_title_slide(prs,
    "《路上有惊慌》",
    "好书推荐 | 作者：饭饭 | 推荐人：四年级X班 XXX")

# 第2页：书籍简介
add_content_slide(prs,
    "📚 这本书讲了什么？",
    [
        "作者饭饭用轻松幽默的笔调，记录生活中的小故事",
        "有家庭的温馨、有社会的观察、还有有趣的思考",
        "书名《路上有惊慌》意思是：人生路上总会遇到意想不到的事情",
        "作者告诉我们：面对惊慌，要保持乐观和勇敢",
        "这是一本适合全家人一起读的随笔集"
    ],
    highlight=3)

# 第3页：精彩片段
add_quote_slide(prs,
    "人生就像走路，不可能一直平坦，也会有坑坑洼洼。",
    "《路上有惊慌》")

# 第4页：推荐理由
add_content_slide(prs,
    "⭐ 我为什么推荐这本书？",
    [
        "语言幽默风趣，读起来轻松愉快",
        "故事贴近生活，让我们学会观察身边的事物",
        "教会我们勇敢面对困难和意外",
        "篇幅短小精悍，适合碎片时间阅读",
        "既有童趣又有哲理，大人小孩都能读"
    ],
    highlight=2)

# 第5页：阅读建议
add_content_slide(prs,
    "📖 阅读小贴士",
    [
        "建议每天读1-2篇，每篇只需几分钟",
        "读到有趣的地方可以和爸爸妈妈分享",
        "试着模仿作者的观察方式，记录自己的生活",
        "遇到不理解的地方不要急，多读几遍",
        "好书值得反复阅读，每次都有新收获"
    ],
    highlight=0)

# 保存
prs.save(OUTPUT_PATH)
print(f"PPT已生成: {OUTPUT_PATH}")
print(f"共 {len(prs.slides)} 页")
